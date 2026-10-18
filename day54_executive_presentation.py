
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# --------------------------------------------------
# Day 54: Executive Business Insights Presentation
# --------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
DATA_FILE = BASE_DIR / "day41_customer_risk_final.csv"
CHART_DIR = BASE_DIR / "day54_visualizations"
OUTPUT_FILE = BASE_DIR / "day54_executive_business_insights.pptx"

df = pd.read_csv(DATA_FILE)

required = [
    "Customer ID",
    "Risk_Level",
    "Retention_Strategy",
    "Total_Sales",
    "Estimated_CLV",
    "Recency_Days",
    "Total_Orders",
]

missing = [column for column in required if column not in df.columns]
if missing:
    raise ValueError(f"Missing dataset columns: {missing}")

for column in [
    "Total_Sales",
    "Estimated_CLV",
    "Recency_Days",
    "Total_Orders",
]:
    df[column] = pd.to_numeric(df[column], errors="coerce")

df = df.dropna(
    subset=[
        "Risk_Level",
        "Total_Sales",
        "Estimated_CLV",
        "Recency_Days",
        "Total_Orders",
    ]
)

# --------------------------------------------------
# Calculate actual business metrics
# --------------------------------------------------

risk_summary = df.groupby("Risk_Level").agg(
    Customers=("Customer ID", "count"),
    Total_Sales=("Total_Sales", "sum"),
    Average_Sales=("Total_Sales", "mean"),
    Estimated_CLV=("Estimated_CLV", "sum"),
    Average_CLV=("Estimated_CLV", "mean"),
    Average_Recency=("Recency_Days", "mean"),
    Average_Orders=("Total_Orders", "mean"),
)

risk_summary["Sales_Share"] = (
    risk_summary["Total_Sales"] / df["Total_Sales"].sum() * 100
)

risk_summary["CLV_Share"] = (
    risk_summary["Estimated_CLV"] / df["Estimated_CLV"].sum() * 100
)

risk_order = ["High Risk", "Medium Risk", "Low Risk"]

total_customers = len(df)
total_sales = df["Total_Sales"].sum()
total_clv = df["Estimated_CLV"].sum()
average_clv = df["Estimated_CLV"].mean()
high_risk_count = int((df["Risk_Level"] == "High Risk").sum())
high_risk_pct = high_risk_count / total_customers * 100

# --------------------------------------------------
# Presentation helpers
# --------------------------------------------------

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(
        Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.7)
    )
    paragraph = title_box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True

    if subtitle:
        sub_box = slide.shapes.add_textbox(
            Inches(0.65), Inches(1.05), Inches(12), Inches(0.45)
        )
        p = sub_box.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(13)

def add_bullets(slide, items, x=0.8, y=1.6, w=11.7, h=4.8, font_size=21):
    box = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()

    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.space_after = Pt(15)

def add_chart_slide(title, filename, takeaway):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title)

    image_path = CHART_DIR / filename
    if image_path.exists():
        slide.shapes.add_picture(
            str(image_path),
            Inches(0.7),
            Inches(1.35),
            width=Inches(7.5),
            height=Inches(4.9),
        )
    else:
        add_bullets(
            slide,
            [f"Chart not found: {filename}"],
            x=0.9,
            y=2.0,
            w=7,
            h=2,
            font_size=18,
        )

    takeaway_box = slide.shapes.add_textbox(
        Inches(8.5), Inches(1.8), Inches(4.1), Inches(4.0)
    )
    tf = takeaway_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "BUSINESS TAKEAWAY"
    p.font.size = Pt(17)
    p.font.bold = True

    p2 = tf.add_paragraph()
    p2.text = takeaway
    p2.font.size = Pt(20)
    p2.space_before = Pt(15)

    return slide

# --------------------------------------------------
# Slide 1: Cover
# --------------------------------------------------

slide = prs.slides.add_slide(prs.slide_layouts[6])

cover_title = slide.shapes.add_textbox(
    Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.4)
)
p = cover_title.text_frame.paragraphs[0]
p.text = "Turning Analytics into Executive Business Insights"
p.font.size = Pt(32)
p.font.bold = True
p.alignment = PP_ALIGN.CENTER

cover_subtitle = slide.shapes.add_textbox(
    Inches(1.3), Inches(3.3), Inches(10.7), Inches(1.1)
)
p = cover_subtitle.text_frame.paragraphs[0]
p.text = "Day 54/60 | ABTalks Data Science Challenge"
p.font.size = Pt(22)
p.alignment = PP_ALIGN.CENTER

cover_note = slide.shapes.add_textbox(
    Inches(1.2), Inches(5.1), Inches(11), Inches(0.7)
)
p = cover_note.text_frame.paragraphs[0]
p.text = "Customer Risk • Sales • Customer Value • Retention"
p.font.size = Pt(17)
p.alignment = PP_ALIGN.CENTER

# --------------------------------------------------
# Slide 2: Executive summary
# --------------------------------------------------

slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Executive Summary", "A snapshot of the customer portfolio")

add_bullets(
    slide,
    [
        f"Customers analysed: {total_customers:,}",
        f"Total sales in dataset: ₹{total_sales:,.2f}",
        f"Total estimated CLV: ₹{total_clv:,.2f}",
        f"Average estimated CLV per customer: ₹{average_clv:,.2f}",
        f"High-risk customers: {high_risk_count:,} ({high_risk_pct:.2f}%)",
    ],
    font_size=22,
)

# --------------------------------------------------
# Slide 3: Risk distribution
# --------------------------------------------------

add_chart_slide(
    "Customer Risk Distribution",
    "day54_risk_distribution.png",
    (
        f"{high_risk_count} customers ({high_risk_pct:.2f}%) "
        "are classified as high risk. Review their behaviour "
        "and retention needs before selecting interventions."
    ),
)

# --------------------------------------------------
# Slide 4: Sales by risk
# --------------------------------------------------

add_chart_slide(
    "Sales Contribution by Risk Group",
    "day54_sales_by_risk.png",
    (
        f"High-risk customers contribute "
        f"{risk_summary.loc['High Risk', 'Sales_Share']:.2f}% "
        "of recorded sales. Assess this segment for possible "
        "re-engagement opportunities."
    ),
)

# --------------------------------------------------
# Slide 5: Recency
# --------------------------------------------------

add_chart_slide(
    "Customer Purchase Recency",
    "day54_recency_by_risk.png",
    (
        f"Average days since last order: High Risk "
        f"{risk_summary.loc['High Risk', 'Average_Recency']:.2f}; "
        f"Medium Risk {risk_summary.loc['Medium Risk', 'Average_Recency']:.2f}; "
        f"Low Risk {risk_summary.loc['Low Risk', 'Average_Recency']:.2f}. "
        "Consider testing targeted engagement for inactive customers."
    ),
)

# --------------------------------------------------
# Slide 6: Retention strategy
# --------------------------------------------------

add_chart_slide(
    "Retention Strategy Distribution",
    "day54_retention_strategies.png",
    (
        "Review how retention strategies are assigned. "
        "The current dataset assigns high-risk customers to "
        "Targeted Re-engagement or Immediate Retention Campaign."
    ),
)

# --------------------------------------------------
# Slide 7: Risk and estimated CLV
# --------------------------------------------------

slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Customer Value at Risk")

high = risk_summary.loc["High Risk"]

add_bullets(
    slide,
    [
        f"High-risk estimated CLV: ₹{high['Estimated_CLV']:,.2f}",
        f"Share of total estimated CLV: {high['CLV_Share']:.2f}%",
        f"Average estimated CLV for high-risk customers: ₹{high['Average_CLV']:,.2f}",
        (
            "Business implication: investigate whether "
            "targeted retention can protect valuable customer relationships."
        ),
        (
            "Important: estimated CLV is a model-derived measure, "
            "not guaranteed future revenue."
        ),
    ],
    font_size=20,
)

# --------------------------------------------------
# Slide 8: Recommendations
# --------------------------------------------------

slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Recommended Business Actions")

add_bullets(
    slide,
    [
        "1. Review high-risk customers with long purchase gaps.",
        "2. Test targeted re-engagement on a defined customer segment.",
        "3. Prioritise outreach using risk, estimated CLV, and business context.",
        "4. Track repeat purchases, recovered sales, and campaign response.",
        "5. Compare results against a baseline before scaling campaigns.",
    ],
    font_size=20,
)

# --------------------------------------------------
# Slide 9: Limitations and next steps
# --------------------------------------------------

slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Limitations and Next Steps")

add_bullets(
    slide,
    [
        "Risk classifications describe model or dataset categories; they do not guarantee churn.",
        "Historical associations do not prove that a retention campaign causes better outcomes.",
        "Estimated CLV should not be interpreted as guaranteed future revenue.",
        "Validate the recommendations with campaign experiments and subsequent customer data.",
        "Monitor customer activity and review results regularly.",
    ],
    font_size=19,
)

# --------------------------------------------------
# Save the presentation
# --------------------------------------------------

prs.save(OUTPUT_FILE)

print("Presentation created successfully!")
print(f"File: {OUTPUT_FILE}")
print(f"Slides: {len(prs.slides)}")
